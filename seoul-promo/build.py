# -*- coding: utf-8 -*-
"""2026 서울 시정 홍보 자료 빌드 스크립트.

content.py 하나를 원본으로 아래 결과물을 만듭니다.
  1) public/seoul-promo/js/slides-data.js   — 웹 3D 프레젠테이션 데이터
  2) seoul-promo/build/*.png                — 3D 장면 이미지 (Playwright 렌더링)
  3) seoul-promo/2026_서울_시정홍보_PPT.pptx  — 20장 PPT (발표자 노트에 나레이션 포함)
  4) seoul-promo/2026_서울_시정홍보_나레이션_대본.txt
  5) seoul-promo/공모전_제출정보.txt

사용법: python3 seoul-promo/build.py [--previews]
필요: python-pptx, Node.js + playwright
"""
import json
import math
import os
import shutil
import subprocess
import sys
import unicodedata

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SITE = os.path.join(ROOT, "public", "seoul-promo")
BUILD = os.path.join(HERE, "build")
sys.path.insert(0, HERE)
import content as C  # noqa: E402

PPTX_NAME = "2026_서울_시정홍보_PPT.pptx"
SCRIPT_NAME = "2026_서울_시정홍보_나레이션_대본.txt"
SUBMIT_NAME = "공모전_제출정보.txt"
WEB_PPTX = "seoul_2026_policy_promo.pptx"

FONT = "맑은 고딕"
NAVY = "0F1F4B"
INK = "16224A"
MUTED = "55617F"
PX = 7620  # 1600px 캔버스 기준 1px = 7620 EMU (13.333in 폭)


# ---------------------------------------------------------------- 1) 웹 데이터
def write_web_data():
    data = {
        "title": C.DECK_TITLE,
        "subtitle": C.DECK_SUBTITLE,
        "parts": C.PARTS,
        "layouts": C.LAYOUTS,
        "slides": [{k: v for k, v in s.items() if k != "narration"} for s in C.SLIDES],
    }
    path = os.path.join(SITE, "js", "slides-data.js")
    with open(path, "w", encoding="utf-8") as f:
        f.write("/* build.py 가 content.py 로부터 생성한 파일입니다. 직접 수정하지 마세요. */\n")
        f.write("window.DECK = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
    print("✓", os.path.relpath(path, ROOT))


# ---------------------------------------------------------------- 2) 3D 이미지
def render_art(previews):
    cmd = ["node", os.path.join(HERE, "render.js"), SITE, BUILD, str(len(C.SLIDES))]
    if previews:
        cmd.append("--previews")
    subprocess.run(cmd, check=True)
    print("✓ 3D 장면 렌더링 →", os.path.relpath(BUILD, ROOT))


# ---------------------------------------------------------------- 3) PPTX
def px(v):
    return Emu(int(round(v * PX)))


def fpt(pxsize):
    """웹 px 글자 크기 → pt (1600px = 13.333in)"""
    return Pt(pxsize * 0.6)


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def mix(h, other, t):
    a = [int(h.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(other.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b))


def text_width(text, size):
    w = 0.0
    for ch in text:
        if ch == " ":
            w += 0.28
        elif unicodedata.east_asian_width(ch) in ("W", "F") or ord(ch) > 0x2E80:
            w += 1.0
        elif ch.isupper() or ch.isdigit():
            w += 0.62
        else:
            w += 0.52
    return w * size


def n_lines(text, size, width):
    total = 0
    for para in text.split("\n"):
        # 단어 단위 줄바꿈 근사
        line, count = 0.0, 1
        for word in para.split(" "):
            ww = text_width(word + " ", size)
            if line + ww > width * 0.94 and line > 0:
                count += 1
                line = ww
            else:
                line += ww
        total += count
    return total


def add_shadow(shape, alpha=12, blur=18, dist=6):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur * 12700), dist=str(dist * 12700),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="16307A")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def set_alpha(shape, alpha):
    sf = shape._element.spPr.find(qn("a:solidFill"))
    clr = sf[0]
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def style_run(run, size, color, bold=False, shadow_color=None):
    f = run.font
    f.size = fpt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    if shadow_color:
        sf = rPr.find(qn("a:solidFill"))
        eff = etree.Element(qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="0", dist=str(4 * 12700), dir="2700000", algn="tl", rotWithShape="0")
        etree.SubElement(sh, qn("a:srgbClr"), val=shadow_color)
        sf.addnext(eff)
    latin = rPr.find(qn("a:latin"))
    ea = etree.Element(qn("a:ea"), typeface=FONT)
    latin.addnext(ea)


def textbox(slide, x, y, w, h, paras, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            spacing=1.2, shadow_color=None, wrap=True):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = paras.split("\n")
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = t
        style_run(r, size, color, bold, shadow_color)
    return tb


def rbox(slide, x, y, w, h, fill, text="", size=24, color=INK, bold=True, line=None, radius=0.18,
         align=PP_ALIGN.LEFT, pad=(20, 12), shadow=True, anchor=MSO_ANCHOR.MIDDLE, dash=False, wrap=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25)
        if dash:
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    if shadow:
        add_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = px(pad[0])
    tf.margin_top = tf.margin_bottom = px(pad[1])
    tf.vertical_anchor = anchor
    if text:
        for i, t in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 1.25
            r = p.add_run()
            r.text = t
            style_run(r, size, color, bold)
    return shp


def background(slide, pc):
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 300
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb("FFFFFF")
    stops[0].position = 0.0
    stops[1].color.rgb = rgb("EEF3FD")
    stops[1].position = 1.0
    # 은은한 장식 원
    for (x, y, d, col, a) in [(1250, -260, 640, pc, 10), (-220, 640, 520, "FFD54A", 16)]:
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(d), px(d))
        o.fill.solid()
        o.fill.fore_color.rgb = rgb(col)
        o.line.fill.background()
        set_alpha(o, a)


def build_pptx():
    prs = Presentation()
    prs.slide_width = px(1600)
    prs.slide_height = px(900)
    blank = prs.slide_layouts[6]
    logo = os.path.join(BUILD, "logo.png")
    total = len(C.SLIDES)

    for i, s in enumerate(C.SLIDES):
        sl = prs.slides.add_slide(blank)
        part = C.PARTS[s["part"]]
        pc = part["color"].lstrip("#")
        lay = C.LAYOUTS[s["layout"]]
        background(sl, pc)

        sl.shapes.add_picture(logo, px(64), px(44), height=px(46))
        if part["id"] > 0:
            label = f"{part['emoji']} {part['name']}"
            w = text_width(label, 20) * 1.1 + 50
            rbox(sl, 1600 - 64 - w, 46, w, 42, pc, label, 20, "FFFFFF", radius=0.5, align=PP_ALIGN.CENTER, pad=(10, 4), wrap=False)

        ax, ay, aw, ah = lay["art"]
        sl.shapes.add_picture(os.path.join(BUILD, f"art-{i + 1:02d}.png"), px(ax), px(ay), px(aw), px(ah))

        tx, ty, tw = lay["text"]
        y = ty
        # 키커
        kw = text_width(s["kicker"], 24) * 1.1 + 40
        rbox(sl, tx, y, kw, 46, mix(pc, "FFFFFF", 0.88), s["kicker"], 24, pc, radius=0.25, pad=(16, 4), shadow=False, wrap=False)
        y += 46 + 22
        # 제목
        tsize = 92 if s["layout"] == "cover" else (58 if s["layout"] == "calendar" else 62)
        widest = max(text_width(t, 1) for t in s["title"].split("\n"))
        tsize = min(tsize, int(tw * 0.95 / widest))
        tl = s["title"].count("\n") + 1
        textbox(sl, tx, y, tw, tl * tsize * 1.22, s["title"], tsize, NAVY, bold=True, spacing=1.08)
        y += tl * tsize * 1.22 + 22
        # 리드
        if s["lead"]:
            ls = 30 if s["layout"] == "cover" else 27
            ll = n_lines(s["lead"], ls, tw)
            textbox(sl, tx, y, tw, ll * ls * 1.5, s["lead"], ls, MUTED, bold=True, spacing=1.3)
            y += ll * ls * 1.5 + 26

        lt = s["layout"]
        if lt == "tips":
            textbox(sl, tx, y, 420, 30, "😥 이런 문제가", 22, MUTED, bold=True)
            textbox(sl, tx + 482, y, 420, 30, "😊 이렇게 해결", 22, MUTED, bold=True)
            y += 40
            cw = (tw - 66) / 2
            for b in s["bullets"]:
                p, q = b.split("|")
                rbox(sl, tx, y, cw, 112, "FFF1E8", p, 24, "A04A12", line="FFB98A", dash=True, shadow=False)
                textbox(sl, tx + cw, y, 66, 112, "➜", 36, pc, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                rbox(sl, tx + cw + 66, y, cw, 112, "E7F8F0", q, 24, "0E6B47", line="8EE0BC")
                y += 112 + 18
        elif lt == "steps":
            cw = (tw - 18) / 2
            for k, b in enumerate(s["bullets"]):
                cx = tx + (k % 2) * (cw + 18)
                cy = y + (k // 2) * (150 + 18)
                rbox(sl, cx, cy, cw, 150, "FFFFFF", b, 25, INK, line="E6ECF7", anchor=MSO_ANCHOR.TOP, pad=(24, 22))
            y += 2 * 168
        elif lt != "calendar" and s["bullets"]:
            size = {"agenda": 32, "prompts": 23}.get(lt, 25)
            padl = {"agenda": 30, "prompts": 70}.get(lt, 56)
            for b in s["bullets"]:
                inner = tw - padl - 24
                bl = n_lines(b, size, inner)
                padv = 22 if lt == "agenda" else (18 if lt == "prompts" else 14)
                h = bl * size * 1.42 + padv * 2
                fill = "F7F5FF" if lt == "prompts" else "FFFFFF"
                line = "E3DBFF" if lt == "prompts" else "E6ECF7"
                bx = rbox(sl, tx, y, tw, h, fill, b, size, INK, line=line, radius=0.14 if h > 80 else 0.25,
                          pad=(0, padv), anchor=MSO_ANCHOR.MIDDLE)
                bx.text_frame.margin_left = px(padl)
                bx.text_frame.margin_right = px(24)
                if lt == "split" or lt == "cover":
                    d = sl.shapes.add_shape(MSO_SHAPE.DIAMOND, px(tx + 18), px(y + h / 2 - 13), px(26), px(26))
                    d.fill.solid()
                    d.fill.fore_color.rgb = rgb(pc)
                    d.line.fill.background()
                elif lt == "prompts":
                    textbox(sl, tx + 18, y + padv - 4, 44, 44, "💬", 30, INK)
                y += h + 14
        if s.get("stat"):
            # 핵심 숫자 카드 — 3D 그림 왼쪽 아래 (웹과 같은 위치)
            st = s["stat"]
            vw = text_width(st["value"], 64) + 16
            uw = text_width(st["unit"], 28) + 10 if st["unit"] else 0
            cw = min(430, max(vw + uw, text_width(st["label"], 20)) + 52)
            ll = n_lines(st["label"], 20, cw - 48)
            ch = 16 + 70 + ll * 27 + 18
            cx, cy = ax, ay + ah - ch
            under = rbox(sl, cx, cy + 10, cw, ch, mix(pc, "E6ECF7", 0.75), shadow=False, radius=0.2)
            card = rbox(sl, cx, cy, cw, ch, "FFFFFF", line="E6ECF7", radius=0.2)
            textbox(sl, cx + 24, cy + 12, vw, 74, st["value"], 64, pc, bold=True, shadow_color=mix(pc, "000000", 0.45), wrap=False)
            if st["unit"]:
                textbox(sl, cx + 24 + vw, cy + 40, uw + 20, 40, st["unit"], 28, NAVY, bold=True, wrap=False)
            textbox(sl, cx + 24, cy + 16 + 70, cw - 48, ll * 27, st["label"], 20, MUTED, bold=True, spacing=1.15)

        # 하단 장식 + 쪽번호
        bar = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(64), px(846), px(160), px(10))
        bar.adjustments[0] = 0.5
        bar.fill.gradient()
        bar.fill.gradient_angle = 0
        bar.fill.gradient_stops[0].color.rgb = rgb(pc)
        bar.fill.gradient_stops[1].color.rgb = rgb("FFD54A")
        bar.line.fill.background()
        textbox(sl, 1300, 836, 236, 30, f"{i + 1:02d} / {total}", 20, "8A95B5", bold=True, align=PP_ALIGN.RIGHT, wrap=False)

        sl.notes_slide.notes_text_frame.text = s["narration"]

    out = os.path.join(HERE, PPTX_NAME)
    prs.save(out)
    os.makedirs(os.path.join(SITE, "downloads"), exist_ok=True)
    shutil.copyfile(out, os.path.join(SITE, "downloads", WEB_PPTX))
    print("✓", os.path.relpath(out, ROOT), "(+ 웹 다운로드용 사본)")


# ---------------------------------------------------------------- 4) 대본 / 제출정보
CPM = 320  # 1분에 읽는 글자 수(공백 포함) 기준


def secs(t):
    return round(len(t) / CPM * 60)


def write_txt(name, lines):
    path = os.path.join(HERE, name)
    with open(path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write("\n".join(lines) + "\n")
    print("✓", os.path.relpath(path, ROOT))


def write_script():
    total = sum(secs(s["narration"]) for s in C.SLIDES)
    L = [
        f"{C.DECK_TITLE} — 슬라이드별 나레이션 대본",
        f"총 {len(C.SLIDES)}장 · 예상 발표 시간 약 {total // 60}분 {total % 60}초 (1분에 약 {CPM}자 기준)",
        "",
        "읽는 요령: 문장 사이 쉼표에서 짧게, 마침표에서 한 박자 쉬어 주세요. 슬라이드가 바뀐 뒤 1초 정도 기다렸다가 읽기 시작하면 자연스럽습니다.",
        "",
    ]
    for i, s in enumerate(C.SLIDES, 1):
        title = s["title"].replace("\n", " ")
        L += [
            "=" * 60,
            f"[{i:02d}/{len(C.SLIDES)}] {s['card']} — {title}",
            f"(약 {secs(s['narration'])}초)",
            "-" * 60,
            s["narration"],
            "",
        ]
    L += ["=" * 60, "참고 자료", *[f"- {x}" for x in C.SOURCES],
          "※ 요금·일정 등은 바뀔 수 있으니 발표 직전 서울시 누리집(seoul.go.kr)에서 한 번 더 확인해 주세요."]
    write_txt(SCRIPT_NAME, L)


def write_submission():
    L = [
        "[AI로 PPT 만들기 공모전] 제출 정보",
        "",
        f"■ 작품명: {C.DECK_TITLE}",
        "■ 공모 분야: AI를 활용해 직접 제작한 PPT (주요 시정 홍보)",
        "■ 제출물",
        f"  1) PPT 파일 원본: {PPTX_NAME} (20장, 발표자 노트에 나레이션 수록)",
        f"  2) 나레이션 대본: {SCRIPT_NAME}",
        "  3) 3D 웹 프레젠테이션: public/seoul-promo/index.html (첫 화면) → slides.html (PPT 보기)",
        "",
        "■ 활용한 AI 도구",
        "  - Claude (Anthropic) / Claude Code: 시정 자료 조사·선별, 20장 구성, 슬라이드 문구, 나레이션 대본 작성,",
        "    CSS 3D 그래픽과 웹페이지 코드, PPTX 자동 생성 스크립트 제작",
        "  ※ 다른 도구(ChatGPT, Canva, Gamma, 미리캔버스 등)를 함께 쓰셨다면 여기에 추가해 적어 주세요.",
        "",
        "■ 핵심 프롬프트(명령어) 예시",
        "  ① 기획",
        "    “서울시 홈페이지를 확인해서 시정 홍보를 하기 좋고 업무 효율에 활용할 수 있는 내용을 찾아 줘.",
        "     누구나 쉽게 이해할 수 있게, 20장 PPT와 장별 나레이션 대본, PPT로 바로 넘어가는 간단한 웹페이지를 만들어 줘.",
        "     PPT 내용은 3D 형식으로 만들어 이해를 도와줘.”",
    ]
    for k, b in enumerate([x for x in C.SLIDES if x["layout"] == "prompts"][0]["bullets"], 2):
        L.append(f"  {'①②③④⑤'[k - 1]} {b}")
    L += [
        "",
        "■ 제작 과정 (4단계)",
        *[f"  {b}" for b in [x for x in C.SLIDES if x["layout"] == "steps"][0]["bullets"]],
        "",
        "■ 시행착오와 해소 방안",
    ]
    for b in [x for x in C.SLIDES if x["layout"] == "tips"][0]["bullets"]:
        p, q = b.split("|")
        L.append(f"  - 문제: {p}\n    해결: {q}")
    L += [
        "  - 문제: 서울시 누리집에 직접 접속이 안 되는 환경이 있었음",
        "    해결: 서울시 공식 보도자료·내 손안에 서울 기사 등을 검색해 교차 확인하고, 불확실한 수치는 싣지 않음",
        "  - 문제: PPT 프로그램에서 3D 표현이 어려움",
        "    해결: 웹(CSS 3D)으로 입체 그림을 그린 뒤 고해상도 이미지로 PPT에 넣고, 글자는 PPT에서 바로 고칠 수 있게 유지",
        "",
        "■ 참고 자료",
        *[f"  - {x}" for x in C.SOURCES],
    ]
    write_txt(SUBMIT_NAME, L)


if __name__ == "__main__":
    write_web_data()
    render_art("--previews" in sys.argv)
    build_pptx()
    write_script()
    write_submission()
